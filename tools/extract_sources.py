"""Extract local, gitignored source texts used by tools/verify_quotes.py.

The official documents are not part of this repo. Download them from 3GPP / IEEE yourself, then run:

    python tools/extract_sources.py \
        --tr-v20   path/to/23700-14-k00.docx \
        --tr-v020  path/to/23700-14-020_MCCclean.docx \
        --ts22137  path/to/22137-l00.docx \
        --tr22870  path/to/22870-k00.docx \
        --tdocs    path/to/SA2_173_tdocs        (folder with the unzipped S2-26xxxxx documents) \
        --chair    path/to/SA2_chair_notes      (folder with the chair-note .doc files; needs antiword) \
        --ieee     path/to/80211bf-2025.pdf     (optional; licensed document)

Output goes to source/ (listed in .gitignore). The paragraph numbers used as ¶ coordinates in the
reports are python-docx paragraph indices of the official TR 23.700-14 v20.0.0 docx, counted over all
paragraphs including empty ones; this script writes them in the first column of source/tr23700-14_v20.txt.
"""
import argparse, glob, html, os, re, subprocess, zipfile

import docx


def docx_paragraphs(path):
    return docx.Document(path).paragraphs


def write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='\n') as f:
        f.write(text)


def docx_text(path):
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.append(' | '.join(c.text.strip() for c in r.cells))
    return '\n'.join(out)


def pptx_text(path):
    z = zipfile.ZipFile(path)
    names = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                   key=lambda n: int(re.findall(r'\d+', n)[0]))
    out = []
    for n in names:
        x = z.read(n).decode('utf-8')
        for p in re.findall(r'<a:p>.*?</a:p>', x, flags=re.S):
            s = ''.join(re.findall(r'<a:t>([^<]*)</a:t>', p))
            if s.strip():
                out.append(html.unescape(s))
    return '\n'.join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--tr-v20', required=True)
    ap.add_argument('--tr-v020')
    ap.add_argument('--ts22137')
    ap.add_argument('--tr22870')
    ap.add_argument('--tdocs')
    ap.add_argument('--chair')
    ap.add_argument('--ieee')
    ap.add_argument('--out', default='source')
    a = ap.parse_args()

    paras = docx_paragraphs(a.tr_v20)
    lines = [f'{i}\t[{p.style.name}]\t{p.text}' for i, p in enumerate(paras) if p.text.strip()]
    write(os.path.join(a.out, 'tr23700-14_v20.txt'), '\n'.join(lines) + '\n')
    d = docx.Document(a.tr_v20)
    rows = [' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows]
    write(os.path.join(a.out, 'tr23700-14_v20_tables.txt'), '\n'.join(rows) + '\n')

    for key, name in (('tr_v020', 'tr23700-14_v020.txt'), ('ts22137', 'ts22137.txt'), ('tr22870', 'tr22870.txt')):
        p = getattr(a, key)
        if p:
            write(os.path.join(a.out, 'extra', name), docx_text(p))

    if a.tdocs:
        for f in glob.glob(os.path.join(a.tdocs, '**', '*.*'), recursive=True):
            tag = os.path.basename(os.path.dirname(f))
            if f.lower().endswith('.docx'):
                write(os.path.join(a.out, 'tdocs', tag + '.txt'), docx_text(f))
            elif f.lower().endswith('.pptx'):
                write(os.path.join(a.out, 'tdocs', tag + '.txt'), pptx_text(f))

    if a.chair:
        for f in glob.glob(os.path.join(a.chair, '**', '*.doc'), recursive=True):
            meeting = os.path.basename(os.path.dirname(f))
            r = subprocess.run(['antiword', '-w', '0', f], capture_output=True, text=True, encoding='utf-8', errors='replace')
            write(os.path.join(a.out, 'chair', f'{meeting}_{os.path.basename(f)[:-4]}.txt'), r.stdout)

    if a.ieee:
        import pypdf
        r = pypdf.PdfReader(a.ieee)
        write(os.path.join(a.out, 'ieee80211bf.txt'), '\n'.join(p.extract_text() or '' for p in r.pages))

    print('source texts written to', a.out)


if __name__ == '__main__':
    main()
